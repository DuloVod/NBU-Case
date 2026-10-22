#!/usr/bin/env python3
"""Build the policy extension without mutating prior data. Offline replay by default."""
from __future__ import annotations
import argparse,csv,hashlib,json,math,re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from urllib.request import Request,urlopen
from urllib.parse import urlencode
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'data/policy/2026-10-22'; RAW=OUT/'raw'
CUTOFF='2026-10-22'; LAST_MONTH='2026-09'
SOURCES={
 'reserves-history.json':('https://bank.gov.ua/ua/tables/int-reserves-dynamics?'+urlencode({'startDate':'01.01.2000','endDate':'01.10.2026'}),'api_json'),
 'reserves-dynamics.html':('https://bank.gov.ua/ua/markets/international-reserves-allinfo/dynamics','source_definition'),
 'fx-page.html':('https://bank.gov.ua/ua/markets/currency-interventions','source_definition'),
 'InterventionsResults.xlsx':('https://bank.gov.ua/files/Finmarket/InterventionsResults.xlsx','weekly_crosscheck_only'),
}
for y in range(2012,2027):
 SOURCES[f'fx-year-{y}.json']=('https://bank.gov.ua/ua/tables/currency-interventions?'+urlencode({'startDate':f'01.01.{y}','endDate':f'31.12.{y}' if y<2026 else '01.10.2026'}),'api_json')

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def dump(name,obj): (OUT/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n')
def write_csv(name,rows,fields=None):
 rows=list(rows); fields=fields or list(rows[0])
 with (OUT/name).open('w',newline='') as f:
  w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)
def month(s): return datetime.strptime(s,'%d.%m.%Y').strftime('%Y-%m')
def download():
 RAW.mkdir(parents=True,exist_ok=True)
 for f,(url,role) in SOURCES.items():
  if (RAW/f).exists():continue
  with urlopen(Request(url,headers={'X-Requested-With':'XMLHttpRequest','Accept':'application/json','User-Agent':'Mozilla/5.0'}),timeout=90) as r:
   b=r.read();(RAW/f).write_bytes(b)
   (RAW/f'{f}.meta.json').write_text(json.dumps({'file':f,'url':url,'status':r.status,'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest(),'last_modified':r.headers.get('Last-Modified'),'retrieved_at':CUTOFF},indent=2))

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--download',action='store_true');args=ap.parse_args()
 OUT.mkdir(parents=True,exist_ok=True)
 if args.download:download()
 existing=list(csv.DictReader((ROOT/'data/inputs/observations_long.csv').open()))
 existing_keys={(r['series'],r['period']) for r in existing}
 # Inventory is recorded BEFORE creating any prepared extension rows.
 inventory={'cutoff':CUTOFF,'raw_existing':[{'file':str(p.relative_to(ROOT)),'bytes':p.stat().st_size,'sha256':sha(p)} for p in sorted((ROOT/'data/raw').iterdir()) if p.is_file()],
  'existing_input_rows':len(existing),'existing_input_series':sorted({r['series'] for r in existing}),
  'existing_auxiliary_monthly_columns':next(csv.reader((ROOT/'data/model/auxiliary_monthly.csv').open())),
  'existing_evidence_note':'Previous policy evidence already includes one September preliminary reserves snapshot; extension claims a new historical series, not discovery of that snapshot.',
  'unchanged_source_policy':'All prior raw/inputs/model and sources are read-only.'}
 dump('inventory_before.json',inventory)
 manifest=[]
 for f,(url,role) in SOURCES.items():
  p=RAW/f
  if not p.exists():raise FileNotFoundError(f'Missing {p}; run --download with network access')
  meta_path=RAW/f'{f}.meta.json';meta=json.loads(meta_path.read_text()) if meta_path.exists() else {}
  manifest.append({'file':str(p.relative_to(ROOT)),'url':url,'role':role,'http_status':meta.get('status',200),'bytes':p.stat().st_size,'sha256':sha(p),'retrieved_at':CUTOFF,'last_modified_http':meta.get('last_modified'),'publication_date':'unknown','publication_note':'HTTP retrieval/Last-Modified is not the original release date.'})
 for f in ['CPI_m.xlsx','CCPI_m.xlsx']:
  p=ROOT/'data/raw'/f
  manifest.append({'file':str(p.relative_to(ROOT)),'url':f'https://bank.gov.ua/files/macro/{f}','role':'existing_raw_reextracted','http_status':None,'bytes':p.stat().st_size,'sha256':sha(p),'retrieved_at':'existing_snapshot_2026-10-20','last_modified_http':None,'publication_date':'unknown','publication_note':'Preserved existing workbook, not a new download.'})
 dump('downloads_manifest.json',manifest)
 rows=[]; defs={}; duplicate_count=0
 def add(series,period,value,units,source_url,raw_path,locator,label,kind,aggregation,date_basis,published_at='unknown',status='observed',increment='new_official_historical_series'):
  if period>LAST_MONTH:return
  if not isinstance(value,(int,float)) or not math.isfinite(value):return
  raw_path=Path(raw_path)
  rows.append({'series':series,'period':period,'value':float(value),'units':units,'status':status,'published_at':published_at,'published_at_status':'known' if published_at!='unknown' else 'unknown','source_url':source_url,'raw_file':str(raw_path.relative_to(ROOT)),'raw_sha256':sha(raw_path),'locator':locator,'retrieved_at':CUTOFF,'information_cutoff':CUTOFF,'date_basis':date_basis,'vintage':'snapshot_2026-10-22_not_historical_release_vintage','increment':increment})
  defs[series]={'series':series,'label_uk':label,'units_monthly':units,'kind':kind,'quarterly_aggregation':aggregation,'date_basis':date_basis,'historical_publication_dates':'unknown unless explicitly recorded','availability_rule_for_backtest':'Use a conservative two-month lag as an ASSUMPTION or reconstruct actual releases; this snapshot is not real-time vintage.'}
 p=RAW/'reserves-history.json';url=SOURCES[p.name][0]
 for i,r in enumerate(json.loads(p.read_text())['results']):
  per=month(r['date']); on=datetime.strptime(r['on_date'],'%d.%m.%Y')
  yr,mo=map(int,per.split('-')); expected=(yr+1,1) if mo==12 else (yr,mo+1)
  assert (on.year,on.month)==expected,('stock date convention',r)
  for key,series,label in [('RES_OffReserveAssets','gross_reserves_usd_mn','Валові міжнародні резерви'),('IRAM_NetReserves','net_reserves_usd_mn','Чисті міжнародні резерви')]:
   if key in r:add(series,per,r[key]*1000,'USD_million',url,p,f'results[{i}].{key}',label,'stock','quarter_end','end_of_reference_month; on_date=first_day_of_next_month')
 # Annual FX API responses overlap preceding December. Compare duplicates, never sum them.
 seen={}
 for p in sorted(RAW.glob('fx-year-*.json')):
  if p.name.endswith('.meta.json'):continue
  url=SOURCES[p.name][0]
  for i,r in enumerate(json.loads(p.read_text())['results']):
   per=month(r['date'])
   assert per==datetime.strptime(r['exp_date'],'%m.%Y').strftime('%Y-%m'),('FX period disagreement',r)
   if per>LAST_MONTH:continue
   for key,series,label,sign in [('VolumNbuICM_BuyPerM','fx_purchases_usd_mn','Купівля валюти НБУ',1),('VolumNbuICM_SelPerM','fx_sales_usd_mn','Продаж валюти НБУ (додатний)',-1),('VolumNbuICM_BalPerM','fx_net_purchases_usd_mn','Сальдо купівля мінус продаж',1)]:
    value=float(r[key])*1000*sign;k=(series,per)
    if k in seen:
     assert abs(seen[k]-value)<1e-7,('conflicting duplicate',k,seen[k],value);duplicate_count+=1;continue
    seen[k]=value
    add(series,per,value,'USD_million',url,p,f'results[{i}].{key}',label,'flow','sum_complete_three_months','transaction_date (date of deal), not settlement/accounting date')
 # Extract twelve national consumption groups, not regional rows or weights.
 groups=['food_non_alcohol','alcohol_tobacco','clothing_footwear','housing_utilities','household_equipment','health','transport','communication','recreation_culture','education','restaurants_hotels','miscellaneous']
 p=ROOT/'data/raw/CPI_m.xlsx';w=load_workbook(p,data_only=True)
 for sn,rate in [('1','yoy'),('2','mom')]:
  s=w[sn]
  for rr,group in enumerate(groups,4):
   for c in range(3,s.max_column+1):
    d=s.cell(2,c).value;value=s.cell(rr,c).value
    if isinstance(d,datetime):
     series=f'cpi_{group}_{rate}';per=d.strftime('%Y-%m')
     assert (series,per) not in existing_keys
     add(series,per,value,'percent_yoy' if rate=='yoy' else 'percent_mom','https://bank.gov.ua/files/macro/CPI_m.xlsx',p,f"'{sn}'!{s.cell(rr,c).coordinate}",str(s.cell(rr,2).value),'rate','quarter_end_yoy' if rate=='yoy' else 'compound_complete_three_months','reference_month; NSA; national group price change, NOT CPI contribution',increment='newly_extracted_existing_raw')
 # The four anchor inflation series are REUSED explicitly for validation/model joining.
 aliases={'CPI_UA_MOM':'cpi_headline_mom','CPI_UA_YOY_PUBLISHED':'cpi_headline_yoy','CORE_CPI_UA_MOM':'cpi_core_mom','CORE_CPI_UA_YOY_PUBLISHED':'cpi_core_yoy'}
 for r in existing:
  if r['series'] not in aliases or r['status']!='observed':continue
  series=aliases[r['series']];rate='mom' if series.endswith('mom') else 'yoy';p=ROOT/'data/raw'/r['source']
  add(series,r['period'],float(r['value']),'percent_mom' if rate=='mom' else 'percent_yoy',f'https://bank.gov.ua/files/macro/{p.name}',p,r['locator'],series,'rate','compound_complete_three_months' if rate=='mom' else 'quarter_end_yoy','reference_month; NSA',increment='reused_existing_prepared_observation')
 rows.sort(key=lambda r:(r['period'],r['series']))
 keys=[(r['series'],r['period']) for r in rows];assert len(keys)==len(set(keys))
 assert all(r['period']<=LAST_MONTH for r in rows)
 assert all(r['published_at']=='unknown' or r['published_at']<=CUTOFF for r in rows)
 wide=defaultdict(dict)
 for r in rows:wide[r['period']][r['series']]=r['value']
 arithmetic=[]
 for per,d in wide.items():
  if 'fx_net_purchases_usd_mn' in d:
   # Source rounding differs by one USD in September 2014. Tolerance: ten USD.
   error=d['fx_purchases_usd_mn']-d['fx_sales_usd_mn']-d['fx_net_purchases_usd_mn'];assert abs(error)<1e-5
   arithmetic.append(abs(error))
  if 'gross_reserves_usd_mn' in d and 'net_reserves_usd_mn'in d:assert d['gross_reserves_usd_mn']>=d['net_reserves_usd_mn']
 assert wide['2026-09']['gross_reserves_usd_mn']==47110.43
 assert abs(wide['2026-09']['fx_net_purchases_usd_mn']+5518.99)<1e-7
 assert abs(wide['2026-09']['cpi_headline_yoy']-10)<1e-7
 fields=['period']+sorted(defs)
 write_csv('long.csv',rows)
 write_csv('monthly.csv',[{'period':p,**d} for p,d in sorted(wide.items())],fields)
 qgroups=defaultdict(list)
 for r in rows:
  y,m=map(int,r['period'].split('-'));qgroups[(f'{y}Q{(m-1)//3+1}',r['series'])].append(r)
 qrows=[];qw=defaultdict(dict)
 for (q,series),rs in sorted(qgroups.items()):
  rs.sort(key=lambda r:r['period']);agg=defs[series]['quarterly_aggregation'];end_month=int(q[-1])*3
  expected={f'{q[:4]}-{m:02d}' for m in range(end_month-2,end_month+1)}
  if {r['period'] for r in rs}!=expected:continue
  if agg.startswith('quarter_end'):val=rs[-1]['value'];unit=rs[-1]['units']
  elif agg.startswith('sum'):val=sum(r['value'] for r in rs);unit='USD_million_per_quarter'
  else:val=100*(math.prod(1+r['value']/100 for r in rs)-1);unit='percent_qoq_NSA'
  qw[q][series]=val;qrows.append({'period':q,'series':series,'value':val,'units':unit,'aggregation':agg,'months_count':3,'source_months':'|'.join(r['period'] for r in rs),'raw_sha256':'|'.join(sorted({r['raw_sha256'] for r in rs})),'status':'derived_from_observed_monthly','published_at':'unknown','vintage':'snapshot_2026-10-22_not_historical_release_vintage'})
 write_csv('quarterly.csv',[{'period':p,**d} for p,d in sorted(qw.items())],fields)
 write_csv('quarterly_long.csv',qrows)
 write_csv('series_dictionary.csv',[defs[k] for k in sorted(defs)])
 coverage=[]
 for series in sorted(defs):
  rr=[r for r in rows if r['series']==series];coverage.append({'series':series,'start':rr[0]['period'],'end':rr[-1]['period'],'observations':len(rr),'increment':rr[0]['increment']})
 write_csv('coverage.csv',coverage)
 for item in inventory['raw_existing']:
  assert sha(ROOT/item['file'])==item['sha256'],('prior raw file changed',item['file'])
 checks={'cutoff':CUTOFF,'last_complete_month':LAST_MONTH,'observations':len(rows),'series':len(defs),'monthly_periods':len(wide),'quarterly_periods':len(qw),'new_official_historical_series_observations':sum(r['increment']=='new_official_historical_series' for r in rows),'newly_extracted_existing_raw_observations':sum(r['increment']=='newly_extracted_existing_raw' for r in rows),'reused_existing_observations':sum(r['increment']=='reused_existing_prepared_observation' for r in rows),'dedup_identical_fx_overlap_removed':duplicate_count,'max_fx_identity_error_usd_mn':max(arithmetic),'dedup_unique':True,'future_observations_excluded':True,'only_complete_quarters':True,'no_historical_release_dates_invented':True,'september_2026':wide['2026-09'],'limitations':['Snapshot has revised data, not real-time historical vintages.','September revised reserves differ from 7 October preliminary release.','Deal-date interventions differ from settlement/accounting amounts.','CPI groups are price changes, not inflation contributions; no weights invented.','No fiscal impulse or external funding series inferred from reserves residual.']}
 dump('checks.json',checks)
 print(json.dumps({k:v for k,v in checks.items() if k not in ['september_2026','limitations']},indent=2))
if __name__=='__main__':main()
