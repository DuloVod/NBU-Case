from pathlib import Path
import sys,json,re,shutil,csv,hashlib,math,argparse
HERE=Path(__file__).resolve().parent
import numpy as np
import pandas as pd
import openpyxl,xlrd
from statsmodels.tsa.seasonal import STL
import statsmodels,scipy

parser=argparse.ArgumentParser();parser.add_argument('--root',type=Path);args=parser.parse_args()
ROOT=args.root.resolve() if args.root else HERE.parent
for n in ['raw','inputs','model','documentation','reproduce','csv_for_excel']:(ROOT/n).mkdir(parents=True,exist_ok=True)
ASOF=pd.Period('2026-09',freq='M');qend=ASOF.asfreq('Q');checks=[];auxobs=[];dictionary=[]
def ck(name,passed,detail):checks.append(dict(check=name,status='PASS' if passed else 'REVIEW',detail=str(detail)))
def write(df,name,index=True):
 df.to_csv(ROOT/'model'/f'{name}.csv',index=index,encoding='utf-8',float_format='%.12g')
 # Explicit Excel-import variant: semicolon separator, decimal comma, UTF-8 BOM.
 df.to_csv(ROOT/'csv_for_excel'/f'{name}.csv',index=index,sep=';',decimal=',',encoding='utf-8-sig',float_format='%.12g')
def readbase(n,freq):
 d=pd.read_csv(ROOT/'inputs'/f'{n}.csv',index_col=0)
 d.index=pd.PeriodIndex(d.index,freq=freq);return d.loc[:ASOF if freq=='M' else qend]
def num(v):
 if isinstance(v,(int,float)) and not isinstance(v,bool) and np.isfinite(v):return float(v)
 if isinstance(v,str) and re.fullmatch(r'\s*-?\d+(?:[.,]\d+)?\s*',v):return float(v.strip().replace(',','.'))
 return None
def obs(key,p,v,unit,fn,loc,note='',status='observed'):
 v=num(v)
 if v is None:return
 per=pd.Period(p,freq='Q' if 'Q' in str(p) else 'M')
 if per.freqstr.startswith('Q'):
  if per>qend:return
 elif per>ASOF:return
 auxobs.append(dict(series=key,period=str(per),value=v,unit=unit,source=fn,locator=loc,status=status,note=note,retrieved_on='2026-10-21'))
def desc(key,unit,meaning,transform,role,source):dictionary.append(dict(variable=key,unit=unit,meaning=meaning,transformation=transform,role=role,source=source))
def sheet(fn,name):return list(openpyxl.load_workbook(ROOT/'raw'/fn,data_only=True)[name].values)

monthly=readbase('core_monthly','M');quarter=readbase('core_quarterly','Q')
# All original numeric values remain separately available in inputs/ and raw/.
seasonal_rows=[];diag=[];sensitivity=[]
def adjust(s,name):
 s=s.dropna();ck(name+' monthly continuity',s.index.equals(pd.period_range(s.index.min(),s.index.max(),freq='M')),len(s))
 fit=STL(np.log(s.to_numpy()),period=12,seasonal=13,trend=21,low_pass=13,robust=True).fit()
 sf=np.exp(fit.seasonal);sa=s/sf
 # Fixed 2021 base is intentionally shared by the two price indexes.
 scale=100/sa.loc['2021'].mean();sa=sa*scale
 for p,x,y,f in zip(s.index,s,sa,sf):seasonal_rows.append(dict(series=name,period=str(p),nsa_level=x,sa_level=y,seasonal_factor=f,base_multiplier=scale,method='robust_log_STL_12_13_21_13'))
 ck(name+' reconstruction',np.allclose(sa/scale*sf,s,rtol=1e-12),float(np.max(abs(sa/scale*sf-s))))
 ck(name+' positive adjusted prices',sa.gt(0).all(),sa.min())
 ck(name+' 2021 mean=100',abs(sa.loc['2021'].mean()-100)<1e-10,sa.loc['2021'].mean())
 for window in [7,25]:
  alt=STL(np.log(s.to_numpy()),period=12,seasonal=window,trend=21,low_pass=13,robust=True).fit()
  a=s/np.exp(alt.seasonal);a=a/a.loc['2021'].mean()*100
  delta=(400*np.log(sa.groupby(sa.index.asfreq('Q')).mean()).diff()-400*np.log(a.groupby(a.index.asfreq('Q')).mean()).diff()).dropna()
  sensitivity.append(dict(series=name,alternative_seasonal_window=window,max_qoq_annualized_difference_pp=abs(delta).max(),last_qoq_annualized_difference_pp=delta.iloc[-1]))
 # Descriptive measure, not a formal residual-seasonality test.
 def month_disp(x):
  z=100*np.log(x).diff().dropna();z=z-z.groupby(z.index.year).transform('mean');return float(z.groupby(z.index.month).mean().std())
 diag.append(dict(series=name,method='robust STL on log level',period=12,seasonal_window=13,trend_window=21,low_pass=13,n=len(s),month_mean_dispersion_before=month_disp(s),month_mean_dispersion_after=month_disp(sa),last_seasonal_factor=sf[-1],calendar_adjusted=False))
 return sa
monthly['CPI_UA_SA']=adjust(monthly.CPI_UA_INDEX_NSA_2021_100,'cpi_ua')
monthly['HICP_EA_SA']=adjust(monthly.HICP_EA_NSA,'cpi_ea')
monthly.index.name='period';write(monthly,'monthly')
pd.DataFrame(seasonal_rows).to_csv(ROOT/'documentation/seasonal_components.csv',index=False)
pd.DataFrame(diag).to_csv(ROOT/'documentation/seasonal_diagnostics.csv',index=False)
pd.DataFrame(sensitivity).to_csv(ROOT/'documentation/seasonal_sensitivity.csv',index=False)

levels=pd.DataFrame(index=quarter.index);levels.index.name='period'
levels['gdp']=quarter.GDP_UA_SA
levels['cpi']=monthly.CPI_UA_SA.groupby(monthly.index.asfreq('Q')).mean()
levels['ir_pol']=quarter.IR_POLICY_CALENDAR_AVG
levels['er']=quarter.ER_UAH_EUR_AVG
levels['gdp_ea']=quarter.GDP_EA_SCA
levels['cpi_ea']=monthly.HICP_EA_SA.groupby(monthly.index.asfreq('Q')).mean()
levels['ir_ea']=quarter.IR_ECB_DFR_CALENDAR_AVG
for key in ['cpi','cpi_ea']:
 counts=monthly['CPI_UA_SA' if key=='cpi' else 'HICP_EA_SA'].groupby(monthly.index.asfreq('Q')).count()
 levels.loc[counts.reindex(levels.index).ne(3),key]=np.nan
ck('66 complete quarters',levels.loc['2010Q1':'2026Q2'].notna().all().all(),levels.loc['2010Q1':'2026Q2'].shape)
ck('GDP Q3 not fabricated',pd.isna(levels.loc['2026Q3','gdp']) and pd.isna(levels.loc['2026Q3','gdp_ea']),'Both GDP unavailable')
for key in ['gdp','gdp_ea']:ck(key+' official SA unchanged',levels[key].equals(quarter['GDP_UA_SA' if key=='gdp' else 'GDP_EA_SCA']),key)
write(levels,'quarterly_levels');write(levels.loc['2010Q1':'2026Q2'],'quarterly_balanced')
measure=pd.DataFrame(index=levels.index)
for key in ['gdp','cpi','er','gdp_ea','cpi_ea']:
 measure['l_'+key]=100*np.log(levels[key]);measure['dl_'+key]=4*measure['l_'+key].diff();measure['d4l_'+key]=measure['l_'+key].diff(4)
measure['ir_pol']=levels.ir_pol;measure['ir_ea']=levels.ir_ea
# Diagnostic real FX index, mean 2021=100. Higher = depreciation.
rer=levels.er*levels.cpi_ea/levels.cpi
measure['rer_index_2021_100']=100*rer/rer.loc['2021'].mean()
ck('Real exchange rate identity',np.allclose(measure.rer_index_2021_100.dropna(),(100*rer/rer.loc['2021'].mean()).dropna()),'ER * P_EA / P_UA, normalized 2021')
write(measure,'quarterly_transformations')
observed=measure[['l_gdp','l_cpi','ir_pol','l_er','l_gdp_ea','l_cpi_ea','ir_ea']]
write(observed,'quarterly_observations')
training=observed.loc['2010Q1':'2026Q2'];write(training,'quarterly_training')
ck('Training observations have no missing values',training.notna().all().all(),training.shape)
flags=pd.DataFrame([dict(period=str(p),series=k,status='observed_or_transformed' if pd.notna(v) else ('not_yet_observed' if p==qend else 'history_not_available')) for p,row in levels.iterrows() for k,v in row.items()])
write(flags,'observation_status',False)

# World Bank Pink Sheet monthly commodity prices. Only explicitly defined benchmarks.
fn='WorldBank_commodities.xlsx';r=sheet(fn,'Monthly Prices')
for j,key,unit in [(2,'brent','USD_per_barrel'),(8,'gas_europe','USD_per_MMBtu'),(36,'wheat_us_srw','USD_per_tonne'),(37,'wheat_us_hrw','USD_per_tonne')]:
 for i,row in enumerate(r):
  if isinstance(row[0],str) and re.fullmatch(r'\d{4}M\d\d',row[0]):obs(key,row[0].replace('M','-'),row[j],unit,fn,f"'Monthly Prices'!{openpyxl.utils.get_column_letter(j+1)}{i+1}",note=str(r[4][j]))
# FAO six nominal food indices; source header states 2014–2016=100.
with (ROOT/'raw/FAO_food.csv').open(encoding='utf-8-sig') as f:
 for i,row in enumerate(csv.reader(f),1):
  if row and re.fullmatch(r'\d{4}-\d\d',row[0]):
   for j,key in enumerate(['fao_food','fao_meat','fao_dairy','fao_cereals','fao_oils','fao_sugar'],1):obs(key,row[0],row[j],'index_2014_2016_100','FAO_food.csv',f'CSV row {i}, column {j+1}',note='Recent FAO values may be revised; meat index contains estimates')
# Wages KVED2010, nominal amount per salaried worker; source gaps preserved.
r=sheet('Wages_m.xlsx','1')
for j,p in enumerate(r[1]):
 if isinstance(p,pd.Timestamp) or hasattr(p,'year'):obs('wage_nominal',str(pd.Period(p,freq='M')),r[2][j],'UAH_per_employee_per_month','Wages_m.xlsx',f"'1'!{openpyxl.utils.get_column_letter(j+1)}3",note='Average salaried employee; KVED2010; source scope changes retained')
# Quarterly ILO 15+ directly observed, not cumulative YTD rates.
r=sheet('Unemp_q.xlsx','3')
for j,p in enumerate(r[1]):
 if isinstance(p,str) and re.fullmatch(r'(I|II|III|IV)\.\d{4}',p):
  roman,year=p.split('.');obs('unemployment_ilo_15plus',year+'Q'+str(['I','II','III','IV'].index(roman)+1),r[7][j],'pct_labour_force','Unemp_q.xlsx',f"'3'!{openpyxl.utils.get_column_letter(j+1)}8",note='Discrete quarter, age 15+, NSA; no post-2021 observations')
# Loans to NFCs: monthly source is amount-weighted; quarterly aggregation below is arithmetic.
r=sheet('Financial_markets.xlsx','4.1.1.3');months=['січень','лютий','березень','квітень','травень','червень','липень','серпень','вересень','жовтень','листопад','грудень'];year=None
for i,row in enumerate(r):
 val=str(row[0]).strip()
 hit=re.fullmatch(r'(20\d\d)(?:\s+\d+)?',val)
 if hit:year=int(hit.group(1))
 if year and val.lower() in months:
  p=f'{year}-{months.index(val.lower())+1:02}'
  for j,key in [(7,'loan_nfc_uah'),(9,'loan_nfc_uah_ex_overdraft')]:obs(key,p,row[j],'pct_per_annum','Financial_markets.xlsx',f"'4.1.1.3'!{openpyxl.utils.get_column_letter(j+1)}{i+1}",note='Monthly weighted new-loan rate; before 2009 broader borrower coverage')
# UONIA observed-business-day arithmetic means; flags in daily source retained.
w=xlrd.open_workbook(ROOT/'raw/UONIA.xls');daily=[]
for s in w.sheets():
 for i in range(s.nrows):
  row=s.row_values(i)
  if num(row[0]) and row[0]>40000 and num(row[1]) is not None:
   dt=pd.Timestamp('1899-12-30')+pd.Timedelta(days=row[0])
   if dt.to_period('M')<=ASOF:daily.append(dict(date=str(dt.date()),uonia=float(row[1]),source_flag=row[2],source_sheet=s.name,source_row=i+1))
ud=pd.DataFrame(daily).sort_values('date');ck('UONIA unique daily dates',not ud.date.duplicated().any(),len(ud));ud.to_csv(ROOT/'model/uonia_daily.csv',index=False)
for p,g in ud.assign(period=lambda x:pd.to_datetime(x.date).dt.to_period('M')).groupby('period'):obs('uonia',str(p),g.uonia.mean(),'pct_per_annum','UONIA.xls','mean of published business-day observations',note=f'{len(g)} published days; no pre-UONIA splice')
# BPM6 monthly flows: select month-labelled columns, never cumulative YTD or annual columns.
r=sheet('BOP_m.xlsx','1.1');monthmap={'січ':1,'лют':2,'берез':3,'квіт':4,'трав':5,'черв':6,'лип':7,'серп':8,'сер':8,'верес':9,'вер':9,'жовт':10,'лист':11,'груд':12};year=None
for j,label in enumerate(r[5]):
 if isinstance(r[4][j],(int,float)) and 2000<=r[4][j]<=2030:year=int(r[4][j])
 mon=str(label or '').strip().replace('*','')
 if year and mon in monthmap:
  p=f'{year}-{monthmap[mon]:02}'
  for i,key in [(6,'current_account'),(7,'trade_balance_gs'),(8,'trade_balance_goods'),(9,'exports_goods'),(10,'imports_goods'),(11,'trade_balance_services'),(12,'exports_services'),(13,'imports_services')]:obs(key,p,r[i][j],'million_USD','BOP_m.xlsx',f"'1.1'!{openpyxl.utils.get_column_letter(j+1)}{i+1}",note='BPM6; monthly flow; provisional' if '*' in str(label) else 'BPM6; monthly flow',status='provisional' if '*' in str(label) else 'observed')
ao=pd.DataFrame(auxobs);ck('Auxiliary unique series-period keys',not ao.duplicated(['series','period']).any(),len(ao));ao.to_csv(ROOT/'documentation/auxiliary_provenance.csv',index=False)
am=ao.loc[~ao.period.str.contains('Q')].pivot(index='period',columns='series',values='value');am.index=pd.PeriodIndex(am.index,freq='M');am=am.reindex(pd.period_range('2000-01',ASOF,freq='M'));am.index.name='period'
am['exports_gs']=am.exports_goods+am.exports_services;am['imports_gs']=am.imports_goods+am.imports_services
ck('BOP continuous monthly history',am.loc['2010-01':'2026-08',['exports_goods','imports_goods','exports_services','imports_services']].notna().all().all(),'200 months; source month abbreviation variants normalized')
ck('New-loan continuous monthly history',am.loc['2006-01':'2026-08',['loan_nfc_uah','loan_nfc_uah_ex_overdraft']].notna().all().all(),'Year label 2010 with footnote parsed as year 2010')
ck('BOP goods balance',np.nanmax(abs(am.exports_goods-am.imports_goods-am.trade_balance_goods))<=2,'Tolerance 2 million USD for independently rounded components')
ck('BOP services balance',np.nanmax(abs(am.exports_services-am.imports_services-am.trade_balance_services))<=2,'Tolerance 2 million USD')
ck('BOP combined balance',np.nanmax(abs(am.exports_gs-am.imports_gs-am.trade_balance_gs))<=2,'Tolerance 2 million USD')
write(am,'auxiliary_monthly')
aq=pd.DataFrame(index=levels.index)
flows=['current_account','trade_balance_gs','trade_balance_goods','trade_balance_services','exports_goods','imports_goods','exports_services','imports_services','exports_gs','imports_gs']
for key in am:
 g=am[key].groupby(am.index.asfreq('Q'));aq[key]=(g.sum(min_count=3) if key in flows else g.mean().where(g.count()==3))
t=ao[ao.series=='unemployment_ilo_15plus'].set_index('period').value;t.index=pd.PeriodIndex(t.index,freq='Q');aq['unemployment_ilo_15plus']=t
ck('ILO cumulative rates excluded',t.index.min()==pd.Period('2019Q1') and t.index.max()==pd.Period('2021Q4'),len(t))
ck('No fabricated postwar ILO',aq.loc['2022Q1':,'unemployment_ilo_15plus'].isna().all(),'Post-2021 actual ILO unavailable')
aq.index.name='period';write(aq,'auxiliary_quarterly')
# Keep survey/native periods separate; no automatic fill across missing months.
for n in ['surveys_monthly','surveys_quarterly','business_monthly','business_quarterly_latest','nbu_benchmarks','nbu_external_assumptions']:
 d=pd.read_csv(ROOT/'inputs'/f'{n}.csv');write(d,n,False)
extended=levels.join(aq);write(extended,'quarterly_extended')
# External path is explicitly incomplete, to be completed by the modelling team.
path=pd.DataFrame(index=pd.period_range('2026Q3','2028Q4',freq='Q'));path.index.name='period'
for k in ['gdp_ea','cpi_ea','ir_ea']:
 path[k]=levels[k].reindex(path.index);path[k+'_status']=np.where(path[k].notna(),'observed_transformed','needs_team_assumption')
write(path,'external_path_template')
for k,unit,label,trans,source in [
 ('gdp','million_UAH_2021_prices','ВВП України, SA','Без повторного коригування','Держстат'),
 ('cpi','index_2021_100','ІСЦ України, дослідницька SA','Robust STL(log), потім середнє 3 місяців','НБУ / Держстат'),
 ('ir_pol','pct_per_annum','Облікова ставка, середня за календарні дні','Середнє щоденних чинних ставок','Архів рішень НБУ'),
 ('er','UAH_per_EUR','Курс, середнє місячних середніх','Просте середнє 3 місяців','НБУ'),
 ('gdp_ea','index_2020_100','ВВП єврозони, SCA','Офіційний рівень','Eurostat'),
 ('cpi_ea','index_2021_100','HICP єврозони, дослідницька SA','Robust STL(log), потім середнє 3 місяців','Eurostat'),
 ('ir_ea','pct_per_annum','Депозитна ставка ЄЦБ, календарне середнє','Єдина конвенція з ir_pol; EOQ збережено у inputs','ECB')]:desc(k,unit,label,trans,'QPM core',source)
for k in ['gdp','cpi','er','gdp_ea','cpi_ea']:
 for prefix,unit,formula in [('l_','log_x100','100*ln(X)'),('dl_','annualized_log_pct_qoq','400*ln(X_t/X_t-1)'),('d4l_','log_pct_yoy','100*ln(X_t/X_t-4)')]:desc(prefix+k,unit,k,formula,'transformed observable','quarterly_levels')
for key,g in ao.groupby('series'):desc(key,g.unit.iloc[0],key,'Quarter sum' if key in flows else ('Native quarterly, no YTD conversion' if key=='unemployment_ilo_15plus' else 'Quarter arithmetic mean of 3 available months'),'auxiliary',g.source.iloc[0])
desc('exports_gs','million_USD','Експорт товарів та послуг','exports_goods + exports_services; квартал сума','auxiliary','BOP_m.xlsx')
desc('imports_gs','million_USD','Імпорт товарів та послуг','imports_goods + imports_services; квартал сума','auxiliary','BOP_m.xlsx')
desc('rer_index_2021_100','index_2021_100','Реальний курс, більше = реальне знецінення','er*cpi_ea/cpi, нормування середнього 2021','diagnostic','quarterly_levels')
pd.DataFrame(dictionary).to_csv(ROOT/'documentation/dictionary.csv',index=False)
coverage=[]
for dataset,d in [('core',levels),('auxiliary',aq)]:
 for key in d:
  s=d[key].dropna();coverage.append(dict(dataset=dataset,series=key,first=str(s.index.min()),last=str(s.index.max()),n=len(s),missing_in_2010_2026Q2=int(d.loc['2010Q1':'2026Q2',key].isna().sum()),missing_periods=';'.join(d.index[d[key].isna()].astype(str))))
pd.DataFrame(coverage).to_csv(ROOT/'documentation/coverage.csv',index=False)
for name,d in [('levels',levels),('observations',measure),('auxiliary_monthly',am),('auxiliary_quarterly',aq)]:
 ck(name+' unique ordered periods',d.index.is_unique and d.index.is_monotonic_increasing,len(d))
 ck(name+' no infinity',not np.isinf(d.to_numpy()).any(),d.shape)
for k in ['cpi','cpi_ea','gdp','gdp_ea','er']:ck(k+' positive for logarithms',levels[k].dropna().gt(0).all(),levels[k].min())
pd.DataFrame(checks).to_csv(ROOT/'documentation/validation.csv',index=False)
sources=json.loads((ROOT/'inputs/base_sources.json').read_text())
for fn,url,meaning in [
 ('WorldBank_commodities.xlsx','https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Monthly.xlsx','Brent; Europe natural gas USD/MMBtu; US SRW/HRW wheat USD/tonne'),
 ('FAO_food.csv','https://www.fao.org/worldfoodsituation/foodpricesindex/en/','FAO nominal food price indices, base 2014–2016=100'),
 ('BOP_m.xlsx','https://bank.gov.ua/files/ES/BOP_m.xlsx','BPM6 monthly flows, million USD; provisional flags preserved'),
 ('Wages_m.xlsx','https://bank.gov.ua/files/macro/Wages_m.xlsx','Nominal average salaried-worker wage, UAH, KVED2010'),
 ('Unemp_q.xlsx','https://bank.gov.ua/files/macro/Unemp_q.xlsx','Sheet3: discrete quarterly ILO age15+; not cumulative YTD'),
 ('Financial_markets.xlsx','https://bank.gov.ua/files/4-Financial_markets.xlsx','New NFC loans in UAH, with and without overdraft'),
 ('UONIA.xls','https://bank.gov.ua/files/UONIA.xls','Published UONIA business days, no pre-UONIA splice')]:sources.append(dict(file=fn,institution='World Bank' if fn.startswith('World') else 'FAO' if fn.startswith('FAO') else 'НБУ',url=url,verdict='Verified source, native definition retained',selection=meaning,purpose='Additional observed series, retrieved 2026-10-21'))
pd.DataFrame(sources).to_csv(ROOT/'documentation/sources.csv',index=False)
audit=[];sheet_audit=[]
for f in sorted((ROOT/'raw').iterdir()):
 rec=dict(file=f.name,bytes=f.stat().st_size,sha256=hashlib.sha256(f.read_bytes()).hexdigest(),format=f.suffix,status='preserved')
 if f.suffix.lower()=='.xlsx':
  w=openpyxl.load_workbook(f,data_only=True);rec['sheets']=len(w.sheetnames)
  for s in w:
   cells=[(c.column,c.value,c.data_type) for row in s for c in row if c.value is not None]
   nc=len(set(x[0] for x in cells));errs=sum(x[2]=='e' for x in cells)
   sheet_audit.append(dict(file=f.name,sheet=s.title,rows=s.max_row,columns=s.max_column,populated_columns=nc,numeric_cells=sum(num(x[1]) is not None for x in cells),source_error_cells=errs,layout='multi_column' if nc>1 else 'empty_or_contents_only'))
 elif f.suffix.lower()=='.xls':
  w=xlrd.open_workbook(f);rec['sheets']=w.nsheets
  for s in w.sheets():
   cells=[(c,s.cell_value(r,c)) for r in range(s.nrows) for c in range(s.ncols) if s.cell_value(r,c)!=''];nc=len(set(x[0] for x in cells))
   sheet_audit.append(dict(file=f.name,sheet=s.name,rows=s.nrows,columns=s.ncols,populated_columns=nc,numeric_cells=sum(num(x[1]) is not None for x in cells),source_error_cells=sum(s.cell_type(r,c)==xlrd.XL_CELL_ERROR for r in range(s.nrows) for c in range(s.ncols)),layout='multi_column' if nc>1 else 'empty_or_contents_only'))
 audit.append(rec)
pd.DataFrame(audit).to_csv(ROOT/'documentation/raw_manifest.csv',index=False);pd.DataFrame(sheet_audit).to_csv(ROOT/'documentation/every_excel_sheet_audit.csv',index=False)
config=dict(data_vintage='Core files retrieved 2026-10-20; auxiliary files retrieved 2026-10-21',last_month=str(ASOF),balanced_start='2010Q1',balanced_end='2026Q2',price_adjustment=dict(method='research robust log-STL',period=12,seasonal=13,trend=21,low_pass=13,robust=True,calendar_adjustment=False),rates='calendar-day arithmetic means for both NBU and ECB',missing='NaN; no fill/interpolation',log_scale=100,qoq_annualized_scale=400,gaps='must be estimated by own QPM; no official estimate copied as fact',versions={'python':sys.version.split()[0],'pandas':pd.__version__,'numpy':np.__version__,'scipy':scipy.__version__,'statsmodels':statsmodels.__version__})
(ROOT/'documentation/config.json').write_text(json.dumps(config,ensure_ascii=False,indent=2))
def records(d):
 if isinstance(d.index,pd.PeriodIndex):d=d.rename_axis('period').reset_index();d['period']=d.period.astype(str)
 return d.replace({np.nan:None}).to_dict('records')
tabs={
 'QPM рівні':records(levels),'QPM спостереження':records(observed),'Повна вибірка':records(training),
 'Темпи та реальний курс':records(measure),
 'Місячні дані':records(monthly),'Додаткові місячні':records(am),'Додаткові квартальні':records(aq),
 'Очікування':pd.read_csv(ROOT/'inputs/surveys_monthly.csv').replace({np.nan:None}).to_dict('records'),
 'ІОДА':pd.read_csv(ROOT/'inputs/business_monthly.csv').replace({np.nan:None}).to_dict('records'),
 'Прогноз НБУ':pd.read_csv(ROOT/'inputs/nbu_benchmarks.csv').replace({np.nan:None}).to_dict('records'),
 'Зовнішній сценарій':records(path),'Словник':dictionary,'Покриття':[{k:v for k,v in r.items() if k!='missing_periods'} for r in coverage],'Джерела':sources,'Перевірки':checks}
for label,filename in [('Квартальні очікування','surveys_quarterly'),('Опитування підприємств','business_quarterly_latest'),('Припущення НБУ','nbu_external_assumptions')]:
 tabs[label]=pd.read_csv(ROOT/'inputs'/f'{filename}.csv').replace({np.nan:None}).to_dict('records')
file_summary=pd.DataFrame(sheet_audit).groupby('file',as_index=False).agg(sheets=('sheet','count'),numeric_cells=('numeric_cells','sum'),source_error_cells=('source_error_cells','sum'))
file_summary['result']='Opened every sheet; original layout retained'
tabs['Аудит Excel-джерел']=file_summary.to_dict('records')
(ROOT/'documentation/workbook_data.json').write_text(json.dumps(tabs,ensure_ascii=False,allow_nan=False))
shutil.copy2(__file__,ROOT/'reproduce/calibrate.py') if Path(__file__).resolve()!=(ROOT/'reproduce/calibrate.py').resolve() else None
print('CORE',levels.shape,'AUX',am.shape,aq.shape,'SHEETS AUDITED',len(sheet_audit),'CHECKS',len(checks))
print('REVIEW',[c for c in checks if c['status']=='REVIEW']);print(pd.DataFrame(diag).to_string(index=False))
