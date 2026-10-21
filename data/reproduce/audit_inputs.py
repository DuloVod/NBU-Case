from pathlib import Path
import sys,json,hashlib,csv
import openpyxl,xlrd
from pypdf import PdfReader
files=['Tables_Charts_2026-Q3.xlsx','Surveys_rate.xlsx','Surveys_price.xlsx','prc_hicp_minr__custom_23410998_spreadsheet.xlsx','PInterestRates.xls','namq_10_gdp__custom_23412352_spreadsheet.xlsx','monthly_long_nf_publ.xlsx','Macroeconomic_Forecast_table_2026_07.xlsx','IR_2026-Q3.pdf','Exchange_r.xls','data.csv','CPI_m.xlsx','CCPI_m.xlsx','BOS_2026-Q3.pdf','_dodatok_sezonka_vvp_2010-2026.xlsx']
packaged=True
root=Path(__file__).resolve().parents[1]
out=root/'reproduce/extracted'
out.mkdir(exist_ok=True,parents=True)
if packaged:
 import shutil
 shutil.copy2(root/'raw/nbu_policy_archive_rows.json',out/'nbu_policy_archive_rows.json')
allrows={};manifest=[]
for fn in files:
 p=root/'raw'/fn
 info={'file':fn,'exists':p.exists()}
 if not p.exists():manifest.append(info);continue
 info.update(bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest())
 if p.suffix in ['.xlsx','.xls']:
  if p.suffix=='.xlsx':
   wb=openpyxl.load_workbook(p,data_only=True)
   sheets={s.title:[list(r) for r in s.iter_rows(values_only=True)] for s in wb}
   wf=openpyxl.load_workbook(p,data_only=False)
   info['formula_count']=sum(c.data_type=='f' for s in wf for row in s for c in row)
   info['errors']=[f'{s.title}!{c.coordinate}: {c.value}' for s in wb for row in s for c in row if c.data_type=='e']
  else:
   wb=xlrd.open_workbook(p)
   sheets={s.name:[s.row_values(r) for r in range(s.nrows)] for s in wb.sheets()}
   info['errors']=[]
  info['sheets']={k:[len(v),max(map(len,v),default=0)] for k,v in sheets.items()}
  allrows[fn]=sheets
  lines=[]
  for name,rows in sheets.items():
   lines.append(f'\nSHEET {name} {len(rows)} rows')
   for i,row in enumerate(rows):
    non=[f'{openpyxl.utils.get_column_letter(j+1)}={str(v)[:180]}' for j,v in enumerate(row) if v is not None and v!='']
    if non:lines.append(f'{i+1}: '+' | '.join(non[:14])+(' ... LAST: '+' | '.join(non[-4:]) if len(non)>14 else ''))
  (out/(fn+'.txt')).write_text('\n'.join(lines))
 elif p.suffix=='.pdf':
  pdf=PdfReader(p);info['pages']=len(pdf.pages)
  (out/(fn+'.txt')).write_text('\n'.join(f'\nPAGE {i+1}\n'+(pg.extract_text()or'') for i,pg in enumerate(pdf.pages)))
 else:
  rows=list(csv.reader(p.open(encoding='utf-8-sig')))
  info['rows']=len(rows);info['header']=rows[0];info['first']=rows[1];info['last']=rows[-1]
 manifest.append(info)
(out/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2,default=str))
(out/'cells.json').write_text(json.dumps(allrows,ensure_ascii=False,default=str))
for x in manifest:print(x['file'],x.get('exists'),x.get('bytes'),len(x.get('errors',[])))
